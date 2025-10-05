import os
from typing import List, Iterator, Optional, Tuple

from docx import Document as DocxDocument
from docx2python import docx2python
from langchain_community.document_loaders.base import BaseLoader
from langchain_core.documents import Document as LCDocument

from rag_helper import DATA_PATH


class DocxByHeadingLoader(BaseLoader):
    """Загрузчик DOCX файлов с разделением по заголовкам и извлечением изображений через docx2python."""

    def __init__(
            self,
            file_path: str,
            heading_style_pattern: str = "Heading",
            images_output_dir: str = "extracted_images",
            extract_images: bool = True
    ):
        """
        Инициализация загрузчика.

        Args:
            file_path: Путь к DOCX файлу
            heading_style_pattern: Паттерн для определения стилей заголовков
            images_output_dir: Директория для сохранения извлеченных изображений
            extract_images: Флаг извлечения изображений
        """
        self.file_path = file_path
        self.heading_style_pattern = heading_style_pattern
        self.images_output_dir = images_output_dir
        self.extract_images = extract_images

        # Создаем директорию для изображений если нужно
        if self.extract_images:
            os.makedirs(self.images_output_dir, exist_ok=True)

    def _is_heading(self, paragraph) -> bool:
        """Проверяет, является ли параграф заголовком на основе стиля."""
        if not paragraph.style or not paragraph.style.name:
            return False
        return self.heading_style_pattern in paragraph.style.name

    def _extract_images_and_text(self) -> Tuple[str, dict]:
        """
        Извлекает изображения и текст с плейсхолдерами, используя docx2python.
        Для каждого документа создается временная папка, чтобы избежать конфликтов.
        """
        try:
            # Создаем уникальное имя временной папки на основе имени документа
            docx_filename = os.path.splitext(os.path.basename(self.file_path))[0]
            temp_image_dir = os.path.join(self.images_output_dir, f"temp_{docx_filename}")

            # Убедимся, что временная папка чистая
            if os.path.exists(temp_image_dir):
                import shutil
                shutil.rmtree(temp_image_dir)
            os.makedirs(temp_image_dir, exist_ok=True)

            # Извлекаем изображения во временную папку
            with docx2python(self.file_path, temp_image_dir) as docx_content:
                text_with_placeholders = docx_content.text

            # Собираем пути К ТОЛЬКО ЧТО ИЗВЛЕЧЕННЫМ изображениям
            image_name_to_path = {}
            temp_image_files = [f for f in os.listdir(temp_image_dir)
                                if f.lower().endswith(('.png', '.jpg', '.jpeg', '.gif', '.bmp'))]

            for img_file in temp_image_files:
                # Создаем новое имя с префиксом имени документа
                new_img_name = f"{docx_filename}_{img_file}"
                old_path = os.path.join(temp_image_dir, img_file)
                new_path = os.path.join(self.images_output_dir, new_img_name)

                # Копируем и переименовываем файл
                import shutil
                shutil.copy2(old_path, new_path)

                # Сохраняем соответствие оригинального имени и нового пути
                image_name_to_path[img_file] = new_path

            # Удаляем временную папку
            import shutil
            shutil.rmtree(temp_image_dir)

            return text_with_placeholders, image_name_to_path

        except Exception as e:
            print(f"Ошибка при извлечении изображений через docx2python: {e}")
            # Убедимся, что временная папка удалена в случае ошибки
            if 'temp_image_dir' in locals() and os.path.exists(temp_image_dir):
                import shutil
                shutil.rmtree(temp_image_dir)
            return "", {}

    def _split_text_with_placeholders_by_headings(self, text_with_placeholders: str, headings: List[str]) -> dict:
        """
        Разделяет текст с плейсхолдерами по заголовкам.

        Args:
            text_with_placeholders: Полный текст с плейсхолдерами
            headings: Список заголовков в порядке их появления

        Returns:
            dict: Словарь {заголовок: текст_секции_с_плейсхолдерами}
        """
        sections = {}
        current_pos = 0

        for i, heading in enumerate(headings):
            # Ищем заголовок в тексте с плейсхолдерами
            start_index = text_with_placeholders.find(heading, current_pos)

            if start_index == -1:
                # Если не нашли, пропускаем эту секцию
                sections[heading] = ""
                continue

            # Находим конец этой секции (начало следующего заголовка или конец текста)
            if i < len(headings) - 1:
                next_heading = headings[i + 1]
                end_index = text_with_placeholders.find(next_heading, start_index + len(heading))
            else:
                end_index = len(text_with_placeholders)

            if end_index == -1:
                end_index = len(text_with_placeholders)

            # Извлекаем текст секции
            section_text = text_with_placeholders[start_index:end_index]
            sections[heading] = section_text
            current_pos = start_index + len(section_text)

        return sections

    def _extract_images_from_section(self, section_text: str, image_name_to_path: dict) -> List[str]:
        """
        Надежный метод извлечения изображений через пошаговый поиск.
        """
        found_images = []

        try:
            # Разбиваем текст на строки для анализа
            lines = section_text.split('\n')

            for line in lines:
                # Ищем строки, содержащие "media/"
                if 'media/' in line:
                    # Извлекаем имя файла после media/
                    parts = line.split('media/')
                    if len(parts) > 1:
                        # Берем часть после media/ и убираем лишние символы
                        file_part = parts[1].split('----')[0] if '----' in parts[1] else parts[1]
                        # Убираем возможные пробелы и спецсимволы
                        file_part = file_part.strip().rstrip('-')

                        # Проверяем, является ли это именем файла изображения
                        if any(file_part.lower().endswith(ext) for ext in ['.png', '.jpg', '.jpeg', '.gif', '.bmp']):
                            if file_part in image_name_to_path:
                                found_images.append(image_name_to_path[file_part])
                            else:
                                # Ищем похожие имена
                                for stored_name, path in image_name_to_path.items():
                                    if stored_name.lower() == file_part.lower():
                                        found_images.append(path)
                                        break
        except Exception as e:
            print(f"Ошибка при пошаговом поиске изображений: {e}")

        return found_images

    def _get_heading_level(self, style_name: Optional[str]) -> Optional[int]:
        """Извлекает уровень заголовка из названия стиля."""
        if not style_name:
            return None

        import re
        match = re.search(r'\d+', style_name)
        return int(match.group()) if match else None

    def lazy_load(self) -> Iterator[LCDocument]:
        """Ленивая загрузка документа с правильным связыванием изображений с секциями."""
        # Извлекаем изображения и текст с плейсхолдерами
        text_with_placeholders, image_name_to_path = "", {}
        if self.extract_images:
            text_with_placeholders, image_name_to_path = self._extract_images_and_text()

        # Загружаем документ для обработки структуры по заголовкам
        doc = DocxDocument(self.file_path)

        # Сначала собираем все заголовки и соответствующие секции
        headings = []
        sections_content = []
        heading_levels = []

        current_heading = "Без заголовка"
        current_heading_level = None
        current_content = []

        for paragraph in doc.paragraphs:
            if self._is_heading(paragraph):
                # Сохраняем предыдущую секцию
                if current_content or current_heading != "Без заголовка":
                    headings.append(current_heading)
                    sections_content.append(current_content)
                    heading_levels.append(current_heading_level)

                # Начинаем новую секцию
                current_heading = paragraph.text
                current_heading_level = self._get_heading_level(paragraph.style.name)
                current_content = []
            else:
                # Добавляем текст параграфа
                if paragraph.text.strip():
                    current_content.append(paragraph.text)

        # Добавляем последнюю секцию
        if current_content or current_heading != "Без заголовка":
            headings.append(current_heading)
            sections_content.append(current_content)
            heading_levels.append(current_heading_level)

        # Разделяем текст с плейсхолдерами по заголовкам
        sections_with_placeholders = {}
        if self.extract_images and text_with_placeholders:
            sections_with_placeholders = self._split_text_with_placeholders_by_headings(
                text_with_placeholders, headings
            )

        # Создаем документы LangChain
        for i, heading in enumerate(headings):
            content = sections_content[i]
            full_text = self._combine_content(heading, content)

            # Находим изображения для этой секции
            section_images = []
            if heading in sections_with_placeholders:
                section_text_with_placeholders = sections_with_placeholders[heading]
                section_images = self._extract_images_from_section(
                    section_text_with_placeholders, image_name_to_path
                )

            metadata = {
                "source": self.file_path,
                "heading": heading,
                "heading_level": heading_levels[i],
                "image_paths": section_images
            }

            yield LCDocument(page_content=full_text, metadata=metadata)

    def _combine_content(self, heading: str, content: List[str]) -> str:
        """Объединяет заголовок и содержание в один текст."""
        if content:
            return f"{heading}\n" + "\n".join(content)
        return heading

    def load(self) -> List[LCDocument]:
        """Загружает все документы сразу."""
        return list(self.lazy_load())


# Пример использования
if __name__ == "__main__":
    # Инициализация загрузчика с извлечением изображений
    loader = DocxByHeadingLoader(
        file_path="C:\\Users\\Yuri\\PycharmProjects\\rag_helper\\data\\rag_data\\Вопросы по пульсу.docx",
        images_output_dir=str(DATA_PATH / "img_storage"),
        extract_images=True
    )

    # Загрузка документов
    documents = loader.load()

    # Вывод результатов
    for i, doc in enumerate(documents):
        print(f"\n--- Документ {i + 1} ---")
        print(f"Заголовок: {doc.metadata['heading']}")
        print(f"Количество изображений в секции: {len(doc.metadata.get('image_paths', []))}")
        print(f"Пути к изображениям: {doc.metadata.get('image_paths', [])}")
        print(f"Содержание: {doc.page_content[:200]}...")
